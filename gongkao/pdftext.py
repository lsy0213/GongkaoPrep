"""PDF / Word 文字提取的公共工具：按版面坐标把文字还原成“视觉行”，去掉页眉页脚和广告行。

真题 PDF 来源杂乱：有的题号排在题干后面、有的选项字母单独成块、每页都有页眉页脚和公众号水印。
这里统一按坐标排序、把同一水平线上的碎片拼成一行，后续解析只和“行”打交道。
"""

import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field

try:
    import pymupdf
    pymupdf.TOOLS.mupdf_display_errors(False)
except ImportError:  # 没装 PyMuPDF 时资料库只能列文件，不能解析
    pymupdf = None

# 资料里夹带的推广文字，出现在解析或材料里时整行去掉
AD_RE = re.compile(
    r"公众号|微信号?[:：]|加微|vx[:：]?|VX|淘宝|闲鱼|店铺|qq群|QQ群|加群|交流群|持续更新|认准|"
    r"首单|包邮|顺丰|扫码|二维码|资料群|考公小镇|樱有尽有|上岸小铺|更多资料|咨询微信|www\.|https?://|\.com\b"
)
QNUM_RE = re.compile(r"^\s*【?\d{1,3}】?\s*[\.．、:：]")
PAGE_NO_RE = re.compile(
    r"^\s*(第\s*\d+\s*页.*|共\s*\d+\s*页|[-—–]?\s*\d+\s*[-—–]?|\d+\s*/\s*\d+|page\s*\d+.*|\d+\s*of\s*\d+)\s*$", re.I
)


@dataclass
class Line:
    page: int
    x0: float
    y0: float
    x1: float
    y1: float
    text: str
    size: float = 0
    bold: bool = False
    tag: str = ""  # header / footer / ad / section …

    @property
    def clean(self):
        return re.sub(r"\s+", "", self.text)


@dataclass
class Doc:
    path: str
    pages: list = field(default_factory=list)  # 每页的 Line 列表（已按阅读顺序排序）
    sizes: list = field(default_factory=list)  # 每页 (宽, 高)
    content_x: tuple = (0, 0)  # 正文左右边界

    def all_lines(self):
        for p in self.pages:
            yield from p


def _merge_spans(spans):
    """把同一水平线上的文字片段合并成行。spans: [(x0,y0,x1,y1,text,size,bold)]"""
    spans = [s for s in spans if s[4].strip()]
    spans.sort(key=lambda s: ((s[1] + s[3]) / 2, s[0]))
    rows = []
    for s in spans:
        cy, h = (s[1] + s[3]) / 2, max(s[3] - s[1], 1)
        for r in rows[-3:]:
            rh = max(r["y1"] - r["y0"], 1)
            if abs(cy - (r["y0"] + r["y1"]) / 2) < min(h, rh) * 0.55:
                r["spans"].append(s)
                r["y0"], r["y1"] = min(r["y0"], s[1]), max(r["y1"], s[3])
                break
        else:
            rows.append({"y0": s[1], "y1": s[3], "spans": [s]})
    out = []
    for r in rows:
        ss = sorted(r["spans"], key=lambda s: s[0])
        text = ""
        prev_x1 = None
        for s in ss:
            t = s[4]
            if prev_x1 is not None and s[0] - prev_x1 > 8 and not text.endswith(" "):
                text += " "
            text += t
            prev_x1 = s[2]
        sizes = [s[5] for s in ss]
        out.append((ss[0][0], r["y0"], max(s[2] for s in ss), r["y1"], text,
                    max(sizes) if sizes else 0, any(s[6] for s in ss)))
    out.sort(key=lambda r: (r[1], r[0]))
    return out


FRAC_L, FRAC_S, FRAC_R = "⦅", "⁄", "⦆"  # 分数写成 ⦅分子⁄分母⦆，阅读器按上下排版显示
BLANK = "____"


def frac_plain(text):
    """⦅24⁄5⦆ → 24/5（搜索、给 AI 看、算答案用）。"""
    return re.sub("⦅([^⁄⦆]*)⁄([^⦆]*)⦆",
                  lambda m: (f"({m.group(1)})" if re.search(r"[+\-×÷]", m.group(1)) else m.group(1)) + "/"
                  + (f"({m.group(2)})" if re.search(r"[+\-×÷]", m.group(2)) else m.group(2)), text)


def _math_pieces(page, spans):
    """竖排分数和答题横线在 PDF 里是“分子、分母两段文字 + 中间一条细横线”和“单独一条细横线”。
    把分子分母并成一个 ⦅分子⁄分母⦆ 片段放在横线高度上，单独的横线变成 ____，这样它们能和同行的“=”排在一起。"""
    try:
        bars = [d["rect"] for d in page.get_drawings()
                if d["rect"].height < 1.6 and 3 < d["rect"].width < 130]
    except Exception:  # noqa: BLE001
        return spans
    if not bars:
        return spans
    hs = sorted(s[3] - s[1] for s in spans) or [10]
    body_h = hs[len(hs) // 2]
    used = set()
    extra = []
    for b in sorted(bars, key=lambda r: (r.y0, r.x0)):
        y = (b.y0 + b.y1) / 2

        def near(above):
            out = []
            for i, s in enumerate(spans):
                if i in used:
                    continue
                cx = (s[0] + s[2]) / 2
                if not (b.x0 - 2 <= cx <= b.x1 + 2) or s[2] - s[0] > (b.x1 - b.x0) + 8:
                    continue
                h = s[3] - s[1]
                if above and y - h * 1.3 <= s[3] <= y + 1.2:
                    out.append(i)
                elif not above and y - 1.2 <= s[1] <= y + h * 0.6:
                    out.append(i)
            return out

        up, down = near(True), near(False)
        if up and down:
            num = "".join(spans[i][4] for i in sorted(up, key=lambda i: spans[i][0]))
            den = "".join(spans[i][4] for i in sorted(down, key=lambda i: spans[i][0]))
            used.update(up + down)
            h = max(spans[i][3] - spans[i][1] for i in up + down)
            extra.append((b.x0, y - h * 0.55, b.x1, y + h * 0.55, f"{FRAC_L}{num}{FRAC_S}{den}{FRAC_R}",
                          max(spans[i][5] for i in up + down), any(spans[i][6] for i in up + down)))
        elif not up and not down and b.width >= 12:
            # 横线正上方压着文字：是给重点词加的下划线，不是答题横线
            over = sum(max(0, min(s[2], b.x1) - max(s[0], b.x0)) for s in spans
                       if y - body_h * 0.9 <= s[3] <= y + 1.5 and s[4].strip())
            if over > b.width * 0.4:
                continue
            # 有文字挨在横线左边或右边同一高度，才算答题横线（排除分隔线、表格线）
            if any(abs(s[3] - y) < body_h * 0.8 and (abs(s[2] - b.x0) < 30 or abs(s[0] - b.x1) < 30) for s in spans):
                extra.append((b.x0, y - body_h * 0.85, b.x1, y + 0.5, BLANK, 0, False))
    return [s for i, s in enumerate(spans) if i not in used] + extra


def page_lines(page, pno, math=False):
    """逐字取坐标：片段里的前导空格会把框撑大，按字重新切块才能排对左右顺序。
    math=True 时还原竖排分数和答题横线（讲义用；真题解析不开）。"""
    return lines_from_pieces(page_pieces(page, math), pno)


def lines_from_pieces(spans, pno):
    return [Line(pno, *r[:5], size=r[5], bold=r[6]) for r in _merge_spans(spans)]


HEAVY_RE = re.compile(r"(?i)bold|heavy|black|semibold|demi|hei|(?<!X)HK|黑|粗")
STRONG_RE = re.compile(r"(?i)bold|heavy|black|semibold|demi|粗")
LIGHT_RE = re.compile(r"(?i)light|thin|xihei|XHK|细")


def heavy_font(name):
    return bool(HEAVY_RE.search(name or "")) and not LIGHT_RE.search(name or "")


def body_font(doc, pages):
    """正文字体：抽样页里字数最多的字体。"""
    from collections import Counter
    c = Counter()
    for i in pages:
        try:
            for b in doc[i].get_text("dict")["blocks"]:
                for l in b.get("lines", []):
                    for s in l["spans"]:
                        c[s.get("font", "")] += len(s["text"].strip())
        except Exception:  # noqa: BLE001
            continue
    return c.most_common(1)[0][0] if c else ""


def emphasized(font, flags, body):
    """加粗：PDF 标了粗体；或者正文是宋体、细黑这类细字体，这一段换成了黑体等粗字体（讲义的小标题常这样）。"""
    if flags & 16:
        return True
    if not font or font == body:
        return False
    if heavy_font(body):
        return bool(STRONG_RE.search(font)) and not STRONG_RE.search(body)
    return heavy_font(font)


def page_pieces(page, math=False, body=None):
    """页面上的文字片段 [(x0, y0, x1, y1, 文字, 字号, 粗体)]，空格和大间距处断开，还没拼成行。
    body：正文字体名，给了就按字体判断加粗。"""
    d = page.get_text("rawdict", flags=pymupdf.TEXT_MEDIABOX_CLIP)
    spans = []
    seen = set()  # 同一个字在同一位置画了两遍（模拟加粗、高亮叠印）：只取一遍，免得“旧石器时代早期旧石器时代早期”
    for b in d["blocks"]:
        if b.get("type") != 0:
            continue
        for l in b["lines"]:
            # 竖排文字（侧边水印）跳过
            if abs(l.get("dir", (1, 0))[1]) > 0.5:
                continue
            for s in l["spans"]:
                size = round(s.get("size", 0), 1)
                bold = emphasized(s.get("font", ""), s.get("flags", 0), body) if body is not None else bool(s.get("flags", 0) & 16)
                piece = []
                for ch in s["chars"]:
                    c = ch["c"]
                    x0, y0, x1, y1 = ch["bbox"]
                    key = (c, round(x0), round(y0))
                    if not c.isspace() and key in seen:
                        continue
                    seen.add(key)
                    if c.isspace():
                        if piece:
                            spans.append(_piece(piece, size, bold))
                            piece = []
                        continue
                    if piece and x0 - piece[-1][2] > max(size, 4) * 1.2:
                        spans.append(_piece(piece, size, bold))
                        piece = []
                    piece.append((x0, y0, x1, y1, c))
                if piece:
                    spans.append(_piece(piece, size, bold))
    spans = _drop_ornaments(spans)
    if math:
        spans = _math_pieces(page, spans)
    return spans


ORNAMENT_RE = re.compile(r"[★☆✩✪✫✬✭✮✯✰⭐◆◇●○■□▲△♦♥♠♣❤❀✿❁]+")


def _drop_ornaments(spans):
    """版式上的装饰符号（晨读讲义每页压着两颗字号 110 的“★”）：字号远大于正文，竖着跨过好几行，
    拼行时会把上下几行搅成一行、顺序打乱，还把“★”粘进标题。只由装饰符号组成、字号超过正文两倍的片段不要。"""
    sizes = sorted(s[5] for s in spans if s[5] and not ORNAMENT_RE.fullmatch(s[4]))
    if not sizes:
        return spans
    body = sizes[len(sizes) // 2]
    return [s for s in spans if not (ORNAMENT_RE.fullmatch(s[4]) and s[5] > body * 2)]


def _piece(chars, size, bold):
    return (min(c[0] for c in chars), min(c[1] for c in chars), max(c[2] for c in chars),
            max(c[3] for c in chars), fix_chars("".join(c[4] for c in chars)), size, bold)


def fix_chars(text):
    """有的 PDF 用康熙部首、兼容汉字代替常用字（⼀→一、⾔→言），统一换回来。"""
    def odd(c):
        o = ord(c)
        return 0x2E80 <= o <= 0x2FDF or 0xF900 <= o <= 0xFAFF or 0x1D400 <= o <= 0x1D7FF  # 末段是数学粗体字母数字 𝟏𝟐

    if not any(odd(c) for c in text):
        return text
    return "".join(unicodedata.normalize("NFKC", c) if odd(c) else c for c in text)


def _norm_hf(text):
    return re.sub(r"\d+", "#", re.sub(r"\s+", "", text))


def mark_header_footer(doc):
    """出现在多页顶部/底部的重复行视为页眉页脚。"""
    n = len(doc.pages)
    cnt = Counter()
    for pi, lines in enumerate(doc.pages):
        h = doc.sizes[pi][1]
        seen = set()
        for ln in lines:
            if (ln.y1 < h * 0.11 or ln.y0 > h * 0.89) and not QNUM_RE.match(ln.text):
                k = _norm_hf(ln.text)
                if k not in seen:
                    cnt[k] += 1
                    seen.add(k)
    need = max(2, int(n * 0.3)) if n > 2 else 99
    for pi, lines in enumerate(doc.pages):
        h = doc.sizes[pi][1]
        for ln in lines:
            top, bottom = ln.y1 < h * 0.11, ln.y0 > h * 0.89
            if (top or bottom) and ((cnt[_norm_hf(ln.text)] >= need and not QNUM_RE.match(ln.text))
                                    or PAGE_NO_RE.match(ln.text)):
                ln.tag = "header" if top else "footer"
            elif AD_RE.search(ln.text) and len(ln.clean) < 60:
                ln.tag = "ad"


def load_pdf(path):
    d = pymupdf.open(path)
    doc = Doc(path)
    for i, page in enumerate(d):
        doc.sizes.append((page.rect.width, page.rect.height))
        doc.pages.append(page_lines(page, i))
    mark_header_footer(doc)
    xs0 = [l.x0 for l in doc.all_lines() if not l.tag]
    xs1 = [l.x1 for l in doc.all_lines() if not l.tag]
    if xs0:
        xs0.sort()
        xs1.sort()
        # 用分位数，避免个别越界的水印把边界撑大
        doc.content_x = (xs0[int(len(xs0) * 0.02)], xs1[int(len(xs1) * 0.98) - 1])
    d.close()
    return doc


def load_docx(path):
    """Word 文档没有版面坐标：每段当一行，页码记为 -1。"""
    import docx

    document = docx.Document(path)
    doc = Doc(path, pages=[[]], sizes=[(0, 0)])
    y = 0
    for para in document.paragraphs:
        for t in para.text.split("\n"):
            if t.strip():
                t = fix_chars(t)
                ln = Line(-1, 0, y, 0, y + 1, t)
                if AD_RE.search(t) and len(ln.clean) < 60:
                    ln.tag = "ad"
                doc.pages[0].append(ln)
                y += 1
    for table in document.tables:
        for row in table.rows:
            t = " ".join(c.text.strip() for c in row.cells if c.text.strip())
            if t:
                doc.pages[0].append(Line(-1, 0, y, 0, y + 1, t))
                y += 1
    return doc


def load_any(path):
    return load_docx(path) if path.lower().endswith(".docx") else load_pdf(path)


def join_text(lines):
    """把若干行拼成段落文字：行尾不是句末标点时直接接上（中文排版的自动换行）。"""
    out = []
    for ln in lines:
        t = ln.text.strip() if isinstance(ln, Line) else str(ln).strip()
        if not t:
            continue
        if out and not re.search(r"[。！？：:；;）)」”\"]$", out[-1]) and not re.match(
            r"^([A-H][\.．、]|[A-D](?=[^\sA-Za-z].*\sB)|[①②③④⑤⑥⑦⑧⑨⑩]|（\d+）|\(\d+\)|\d+[\.．、]|[一二三四五六七八九十]+、)", t
        ):
            out[-1] += (" " if re.match(r"^[A-D](?![A-Za-z])", t) else "") + t
        else:
            out.append(t)
    return "\n".join(out)
